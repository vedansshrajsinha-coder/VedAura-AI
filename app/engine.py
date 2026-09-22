import io

from app.memory_handler import is_memory_message, remember, recall_memory
from app.commands import handle_command
from brain_gateway import ask as brain_ask


def _document_evidence(uploaded_files):
    """
    Convert supported uploaded documents into temporary evidence.
    The files are NOT inserted into the persistent shared knowledge DB.
    """
    evidence = []

    if not uploaded_files:
        return evidence

    from pypdf import PdfReader
    from openpyxl import load_workbook

    for uploaded_file in uploaded_files:
        filename = (uploaded_file.filename or "").strip()
        lower = filename.lower()

        try:
            raw = uploaded_file.read()
            text = ""

            if lower.endswith((".txt", ".md")):
                text = raw.decode("utf-8", errors="ignore")

            elif lower.endswith(".pdf"):
                reader = PdfReader(io.BytesIO(raw))
                text = "\n\n".join(
                    f"[PAGE {i + 1}]\n{page.extract_text() or ''}"
                    for i, page in enumerate(reader.pages)
                )

            elif lower.endswith(".xlsx"):
                workbook = load_workbook(
                    io.BytesIO(raw),
                    read_only=True,
                    data_only=True
                )
                lines = []

                for sheet in workbook.worksheets:
                    lines.append(f"[SHEET {sheet.title}]")

                    for row in sheet.iter_rows(values_only=True):
                        lines.append(
                            " | ".join(
                                "" if cell is None else str(cell)
                                for cell in row
                            )
                        )

                text = "\n".join(lines)

            elif lower.endswith((".html", ".htm")):
                from bs4 import BeautifulSoup

                soup = BeautifulSoup(
                    raw.decode("utf-8", errors="ignore"),
                    "html.parser"
                )

                for node in soup(["script", "style", "nav"]):
                    node.decompose()

                text = soup.get_text("\n")

            elif lower.endswith(".docx"):
                from docx import Document

                document = Document(io.BytesIO(raw))
                text = "\n".join(
                    paragraph.text
                    for paragraph in document.paragraphs
                )

            elif lower.endswith(".epub"):
                import os
                import tempfile

                from ebooklib import epub, ITEM_DOCUMENT
                from bs4 import BeautifulSoup

                with tempfile.NamedTemporaryFile(
                    suffix=".epub",
                    delete=False
                ) as tmp:
                    tmp.write(raw)
                    tmp_path = tmp.name

                try:
                    book = epub.read_epub(tmp_path)
                    text = "\n".join(
                        BeautifulSoup(
                            item.get_content(),
                            "html.parser"
                        ).get_text("\n")
                        for item in book.get_items_of_type(ITEM_DOCUMENT)
                    )
                finally:
                    os.unlink(tmp_path)

            text = text.strip()

            if text:
                evidence.append({
                    "title": filename or "Uploaded document",
                    "source_path": f"upload:{filename}",
                    "text": text[:18000]
                })

        except Exception as exc:
            evidence.append({
                "title": filename or "Uploaded file",
                "source_path": f"upload:{filename}",
                "text": f"File could not be parsed: {exc}"
            })

    return evidence


def process_message(
    user_message,
    uploaded_files=None,
    conversation_history=None
):
    # Existing VedAura commands remain available.
    command_reply = handle_command(user_message)
    if command_reply:
        return command_reply

    # Existing explicit memory commands remain available.
    if is_memory_message(user_message):
        return remember(user_message)

    memory_reply = recall_memory(user_message)
    if memory_reply:
        return memory_reply

    # Normal conversational requests now go through AI Brain.
    result = brain_ask(
        user_message,
        history=conversation_history or [],
        extra_evidence=_document_evidence(uploaded_files)
    )

    return result["answer"]
